import os
import re
from openpyxl import Workbook
import PyPDF2
from docx import Document

def extract_info(text):
  """Extracts email IDs, contact numbers, and overall text from CV text."""
  email_regex = r"[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+"
  phone_regex = r"\d{3}-\d{3}-\d{4}|\d{10}"
  email = re.search(email_regex, text)
  phone_number = re.search(phone_regex, text)
  return {
      "email": email.group() if email else "",
      "phone_number": phone_number.group() if phone_number else "",
      "text": text
  }

def extract_text_from_pdf(pdf_file):
  """Extracts text from a PDF file."""
  try:
    with open(pdf_file, 'rb') as pdf:
      pdf_reader = PyPDF2.PdfReader(pdf)
      text = ""
      for page_num in range(len(pdf_reader.pages)):
        text += pdf_reader.pages[page_num].extract_text()
      return text
  except Exception as e:
    print(f"Error processing PDF {pdf_file}: {e}")
    return ""

def extract_text_from_docx(docx_file):
  """Extracts text from a DOCX file."""
  try:
    doc = Document(docx_file)
    text = ""
    for paragraph in doc.paragraphs:
      text += paragraph.text
    return text
  except Exception as e:
    print(f"Error processing DOCX {docx_file}: {e}")
    return ""

def process_cvs(cv_folder, output_filename):
  """Processes CVs in a folder and saves the extracted information in .xlsx format."""
  wb = Workbook()
  ws = wb.active
  ws.append(["Email", "Phone Number", "Text"])  # Header row

  for root, _, files in os.walk(cv_folder):
    for filename in files:
      if filename.endswith(".pdf"):
        text = extract_text_from_pdf(os.path.join(root, filename))
      elif filename.endswith(".docx"):
        text = extract_text_from_docx(os.path.join(root, filename))
      else:
        print(f"Skipping unsupported file: {filename}")
        continue

      info = extract_info(text)
      ws.append([info["email"], info["phone_number"], info["text"]])

  wb.save(output_filename)  # Save the Excel file

if __name__ == "__main__":
  cv_folder = r"D:\Sample2_2024\Sample2"  # Replace with your folder path
  output_filename = "extracted_info.xlsx"
  process_cvs(cv_folder, output_filename)
